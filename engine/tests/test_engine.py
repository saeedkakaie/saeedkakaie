import io

import numpy as np
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from ibp_engine import portfolio
from ibp_engine.app import app

PERIODS = [1403 * 100 + m for m in range(8, 13)] + [1404 * 100 + m for m in range(1, 13)] + [1405 * 100 + m for m in range(1, 8)]
SEASON = np.array([0.9, 1.0, 1.2, 1.4, 1.4, 1.2, 1.0, 0.8, 0.7, 0.7, 0.8, 0.9])


def _series(base, growth, noise, seed, season=True):
    rng = np.random.default_rng(seed)
    t = np.arange(24)
    months = np.array([p % 100 for p in PERIODS]) - 1
    s = SEASON[months] if season else 1.0
    return list(np.maximum(0, base * (1 + growth) ** (t / 12) * s * (1 + noise * rng.standard_normal(24))))


def _payload():
    a = _series(1000, 0.3, 0.03, 1)
    b = _series(800, -0.25, 0.03, 2)
    c = _series(500, 0.0, 0.25, 3, season=False)
    sporadic = [0 if i % 2 else 300 for i in range(24)]
    return {
        "periods": PERIODS,
        "products": [
            {"id": "a", "familyId": "f1", "revenue": 9e9, "margin": 0.35, "growth": 0.3, "cv": 0.25, "qty": a},
            {"id": "b", "familyId": "f1", "revenue": 6e9, "margin": 0.2, "growth": -0.25, "cv": 0.2, "qty": b},
            {"id": "c", "familyId": "f2", "revenue": 2e9, "margin": 0.3, "growth": 0.0, "cv": 0.3, "qty": c},
            {"id": "d", "familyId": "f2", "revenue": 5e8, "margin": 0.1, "growth": 0.05, "cv": 1.0, "qty": sporadic},
        ],
    }


def test_decompose_detects_trend_and_season():
    out = portfolio.analyze(_payload())
    by = {p["id"]: p for p in out["products"]}
    assert by["a"]["trendDir"] == "up" and by["a"]["seasonalStrength"] > 0.6
    assert by["c"]["seasonalStrength"] < 0.3  # نوسان تصادفی نباید فصلی تشخیص داده شود
    assert by["b"]["trendDir"] == "down"
    assert by["d"]["seasonalStrength"] is None  # سری متناوب تجزیه نمی‌شود
    assert by["a"]["peakMonth"] in (4, 5)


def test_segments_and_concentration():
    out = portfolio.analyze(_payload())
    assert len(out["clusters"]) == 4
    assert {p["cluster"] for p in out["products"]} == {c["id"] for c in out["clusters"]}
    assert 0 < out["concentration"]["gini"] < 1
    assert out["concentration"]["hhi"] > 2500


def test_substitution_pair_detected():
    rng = np.random.default_rng(7)
    base = np.array(_series(1000, 0.0, 0.0, 0))
    eps = rng.standard_normal(24) * 0.2
    payload = {"periods": PERIODS, "products": [
        {"id": "x", "familyId": "f", "revenue": 1, "qty": list(base * (1 + eps))},
        {"id": "y", "familyId": "f", "revenue": 1, "qty": list(base * (1 - eps))},
    ]}
    pairs = portfolio.analyze(payload)["pairs"]
    assert pairs and pairs[0]["r"] < -0.8


def test_http_endpoints_and_report():
    client = TestClient(app)
    h = client.get("/health").json()
    assert h["status"] == "ok" and "portfolio.analyze" in h["capabilities"]
    r = client.post("/portfolio/analyze", json=_payload())
    assert r.status_code == 200 and len(r.json()["products"]) == 4
    spec = {"title": "t", "sheets": [{"name": "پورتفوی", "columns": [{"header": "محصول"}, {"header": "فروش", "type": "money"}, {"header": "سهم", "type": "pct"}],
                                      "rows": [["الف", 100, 0.5], ["ب", 60, 0.3]],
                                      "charts": [{"type": "bar", "x": 0, "y": [1], "title": "فروش"}, {"type": "pie", "x": 0, "y": [2]}]}]}
    x = client.post("/reports/xlsx", json=spec)
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    ws = wb["پورتفوی"]
    assert ws.sheet_view.rightToLeft and len(ws._charts) == 2 and ws.cell(5, 1).value == "الف"
    # بخش drawing با پیشوند استاندارد Excel تا برنامه (ExcelJS) هم بتواند فایل را بخواند
    import zipfile
    drawing = zipfile.ZipFile(io.BytesIO(x.content)).read("xl/drawings/drawing1.xml")
    assert b"<xdr:wsDr" in drawing
