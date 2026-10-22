"""گزارش Excel تحلیلی با نمودارهای بومی Excel (openpyxl).

برنامه جدول‌ها را می‌فرستد و موتور یک کتاب کار راست‌به‌چپ با قالب‌بندی و نمودار برمی‌گرداند.
ساختار ورودی:
{
  "title": "...", "subtitle": "...",
  "sheets": [
    {"name": "...", "title": "...", "columns": [{"header": "...", "type": "money|pct|int|num|text"}],
     "rows": [[...], ...],
     "charts": [{"type": "bar|line|pie|scatter", "title": "...", "x": 0, "y": [2, 3], "anchor": "K4",
                 "xTitle": "...", "yTitle": "...", "limit": 20}]}
  ]
}
"""
from __future__ import annotations

import io
import re
import xml.etree.ElementTree as ET
import zipfile
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference, ScatterChart, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Tahoma"
HEADER_FILL = PatternFill("solid", fgColor="1E3A8A")
ZEBRA_FILL = PatternFill("solid", fgColor="F5F7FB")
THIN = Side(style="thin", color="E4E8F0")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
NUMFMT = {"money": "#,##0", "int": "#,##0", "num": "#,##0.0", "dec": "#,##0.00", "pct": "0.0%"}
PALETTE = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "6250D6", "E34948"]
HEADER_ROW = 4


def _sheet_title(name: str, used: set[str]) -> str:
    base = "".join("-" if ch in '[]:*?/\\' else ch for ch in (name or "Sheet"))[:31] or "Sheet"
    title, i = base, 2
    while title in used:
        suffix = f" ({i})"
        title = base[: 31 - len(suffix)] + suffix
        i += 1
    used.add(title)
    return title


def _write_table(ws, spec: dict[str, Any], subtitle: str) -> int:
    cols = spec.get("columns", [])
    rows = spec.get("rows", [])
    ncol = max(len(cols), 2)
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws.cell(1, 1, spec.get("title") or spec.get("name")).font = Font(name=FONT, bold=True, size=14, color="0B1220")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    ws.cell(2, 1, subtitle).font = Font(name=FONT, size=9, color="667085")
    widths = []
    for ci, col in enumerate(cols, start=1):
        c = ws.cell(HEADER_ROW, ci, col.get("header"))
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        widths.append(max(10, len(str(col.get("header", ""))) + 2))
    ws.row_dimensions[HEADER_ROW].height = 30
    for ri, row in enumerate(rows):
        r = HEADER_ROW + 1 + ri
        for ci, col in enumerate(cols, start=1):
            v = row[ci - 1] if ci - 1 < len(row) else None
            c = ws.cell(r, ci, v)
            c.font = Font(name=FONT, size=10)
            c.border = BORDER
            fmt = NUMFMT.get(col.get("type", ""))
            if fmt:
                c.number_format = fmt
            if ri % 2 == 1:
                c.fill = ZEBRA_FILL
            if v is not None:
                widths[ci - 1] = max(widths[ci - 1], min(48, len(str(v)) + 2))
    for ci, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = ws.cell(HEADER_ROW + 1, 1)
    if rows:
        ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(cols))}{HEADER_ROW + len(rows)}"
    return len(rows)


def _add_chart(ws, chart_spec: dict[str, Any], nrows: int) -> None:
    if nrows == 0:
        return
    kind = chart_spec.get("type", "bar")
    limit = min(nrows, int(chart_spec.get("limit") or nrows))
    first, last = HEADER_ROW + 1, HEADER_ROW + limit
    x_col = int(chart_spec.get("x", 0)) + 1
    y_cols = [int(y) + 1 for y in chart_spec.get("y", [1])]
    cats = Reference(ws, min_col=x_col, min_row=first, max_row=last)
    if kind == "scatter":
        chart = ScatterChart()
        chart.style = 13
        for i, yc in enumerate(y_cols):
            vals = Reference(ws, min_col=yc, min_row=HEADER_ROW, max_row=last)
            s = Series(vals, Reference(ws, min_col=x_col, min_row=first, max_row=last), title_from_data=True)
            s.marker.symbol = "circle"
            s.marker.size = 8
            s.marker.graphicalProperties.solidFill = PALETTE[i % len(PALETTE)]
            s.marker.graphicalProperties.line.solidFill = PALETTE[i % len(PALETTE)]
            s.graphicalProperties.line.noFill = True
            chart.series.append(s)
    elif kind == "pie":
        chart = PieChart()
        chart.add_data(Reference(ws, min_col=y_cols[0], min_row=HEADER_ROW, max_row=last), titles_from_data=True)
        chart.set_categories(cats)
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
    else:
        chart = LineChart() if kind == "line" else BarChart()
        if kind == "bar":
            chart.type = chart_spec.get("direction", "col")
            chart.gapWidth = 60
        for yc in y_cols:
            chart.add_data(Reference(ws, min_col=yc, min_row=HEADER_ROW, max_row=last), titles_from_data=True)
        chart.set_categories(cats)
        for i, s in enumerate(chart.series):
            color = PALETTE[i % len(PALETTE)]
            if kind == "line":
                s.graphicalProperties.line.solidFill = color
                s.graphicalProperties.line.width = 22000
                s.smooth = False
            else:
                s.graphicalProperties.solidFill = color
                s.graphicalProperties.line.solidFill = color
    chart.title = chart_spec.get("title")
    if hasattr(chart, "x_axis") and chart_spec.get("xTitle"):
        chart.x_axis.title = chart_spec["xTitle"]
    if hasattr(chart, "y_axis") and chart_spec.get("yTitle"):
        chart.y_axis.title = chart_spec["yTitle"]
    if hasattr(chart, "x_axis"):
        chart.x_axis.delete = False
        chart.y_axis.delete = False
    chart.height = float(chart_spec.get("height", 9))
    chart.width = float(chart_spec.get("width", 18))
    ws.add_chart(chart, chart_spec.get("anchor") or f"{get_column_letter(len(ws[HEADER_ROW]) + 2)}{HEADER_ROW}")


DRAWING_NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
}


def _normalize_drawings(data: bytes) -> bytes:
    """بازنویسی بخش‌های drawing با پیشوندهای استاندارد Excel (xdr:، a:).

    openpyxl این بخش‌ها را با فضای نام پیش‌فرض می‌نویسد؛ Excel آن را می‌پذیرد اما برخی کتابخانه‌ها
    (از جمله ExcelJS که برنامه برای خواندن فایل‌ها به کار می‌برد) فقط شکل پیشونددار را می‌شناسند.
    """
    for prefix, uri in DRAWING_NS.items():
        ET.register_namespace(prefix, uri)
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            content = src.read(item.filename)
            if re.match(r"xl/drawings/drawing\d+\.xml$", item.filename):
                root = ET.fromstring(content)
                content = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' + ET.tostring(root, encoding="utf-8")
            dst.writestr(item, content)
    return out.getvalue()


def build_report(spec: dict[str, Any]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    subtitle = " • ".join(x for x in [spec.get("company"), spec.get("subtitle"), spec.get("date")] if x)
    for sh in spec.get("sheets", []):
        ws = wb.create_sheet(_sheet_title(sh.get("name", "Sheet"), used))
        n = _write_table(ws, sh, subtitle)
        for ch in sh.get("charts", []):
            _add_chart(ws, ch, n)
    if not wb.worksheets:
        wb.create_sheet("خالی")
    buf = io.BytesIO()
    wb.save(buf)
    return _normalize_drawings(buf.getvalue())
