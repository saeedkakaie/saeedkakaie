# -*- coding: utf-8 -*-
"""خروجی اکسل حرفه‌ای گزارش عارضه‌یابی بانک کالای فعلی (بر پایه‌ی write_report اصلی)."""
import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill('solid', fgColor='0F6E56')
HEADER_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=10)
CELL_FONT = Font(name='Arial', size=10)
TITLE_FONT = Font(name='Arial', bold=True, size=14, color='0F6E56')
thin = Side(style='thin', color='DDDDDD')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def _write_sheet(wb, name, df):
    ws = wb.create_sheet(name[:31])
    ws.sheet_view.rightToLeft = True
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=1, column=j, value=col)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal='center', wrap_text=True)
        c.border = BORDER
    for row in df.itertuples(index=False):
        ws.append(list(row))
    for i in range(2, len(df) + 2):
        for j in range(1, len(df.columns) + 1):
            cell = ws.cell(row=i, column=j)
            cell.font = CELL_FONT
            cell.border = BORDER
            cell.alignment = Alignment(horizontal='right', vertical='center')
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 28
    for j in range(1, len(df.columns) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 24
    return ws


def build_report_xlsx(kpi: dict, clusters: list, records: list) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet('خلاصه')
    ws.sheet_view.rightToLeft = True
    ws['B2'] = 'گزارش عارضه‌یابی و پاکسازی بانک کالا'
    ws['B2'].font = TITLE_FONT
    r = 4
    for k, v in kpi.items():
        ws.cell(row=r, column=2, value=k).font = CELL_FONT
        ws.cell(row=r, column=3, value=v).font = Font(name='Arial', bold=True)
        r += 1
    ws.column_dimensions['B'].width = 44
    ws.column_dimensions['C'].width = 16

    clusters_df = pd.DataFrame([{
        'نوع خوشه': 'تکرار دقیق' if c['kind'] == 'exact' else 'تکرار نرم (فرمتی)',
        'شرح نرمال‌شده': c['normalized_desc'],
        'تعداد کد': len(c['codes']),
        'کدها': ', '.join(str(x) for x in c['codes']),
        'کد پیشنهادی (طلایی)': c['golden_code'],
        'کدهای منسوخ‌شونده': ', '.join(str(x) for x in c['codes'] if str(x) != str(c['golden_code'])),
        'ناهماهنگی گروه؟': 'بله' if c['group_conflict'] else 'خیر',
        'ناهماهنگی واحد؟': 'بله' if c['unit_conflict'] else 'خیر',
    } for c in clusters])
    if len(clusters_df):
        _write_sheet(wb, 'خوشه‌های تکراری', clusters_df)
    else:
        _write_sheet(wb, 'خوشه‌های تکراری', pd.DataFrame([{'نتیجه': 'خوشه‌ی تکراری یافت نشد'}]))

    rec_df = pd.DataFrame([{
        'پلنت': r_['plant'],
        'کد کالا': r_['code'],
        'گروه خام': r_['material_group_raw'],
        'گروه استاندارد پیشنهادی': r_['material_group_mapped'],
        'شرح کالا (خام)': r_['description_raw'],
        'شرح کالا (نرمال‌شده پیشنهادی)': r_['description_normalized'],
        'واحد': r_['unit'],
        'کد نهایی پیشنهادی': r_['suggested_final_code'],
        'خرابی نقل‌قول؟': 'بله' if r_['quote_issue'] else 'خیر',
        'نیازمند نرمال‌سازی متن؟': 'بله' if r_['needs_text_normalization'] else 'خیر',
        'گروه خالی؟': 'بله' if r_['group_missing'] else 'خیر',
        'ناهماهنگی واحد؟': 'بله' if r_['unit_conflict'] else 'خیر',
        'تکرار دقیق؟': 'بله' if r_['duplicate_exact'] else 'خیر',
        'تکرار نرم؟': 'بله' if r_['duplicate_loose'] else 'خیر',
        'منطبق با استاندارد جدید؟': 'خیر' if r_['non_compliant_new_rule'] else 'بله',
        'دلایل عدم تطابق': ' | '.join(r_['non_compliance_reasons']),
    } for r_ in records])
    _write_sheet(wb, 'داده کامل پاکسازی‌شده', rec_df)

    noncompliant_df = rec_df[rec_df['منطبق با استاندارد جدید؟'] == 'خیر']
    _write_sheet(wb, 'ناسازگار با استاندارد جدید', noncompliant_df if len(noncompliant_df) else pd.DataFrame([{'نتیجه': 'موردی یافت نشد'}]))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
